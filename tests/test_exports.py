"""export_reports.py: all five Tableau CSVs plus the Excel tracker are
produced with the expected columns, snapshot stamp, and non-null keys.

The export script is pointed at the throwaway test schema via the
FABTOOL_DB_SCHEMA env var and a fixed --as-of, so expectations are
deterministic (the fixture).
"""

import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import export_reports  # noqa: E402

from conftest import SCHEMA, AS_OF  # noqa: E402

CSV_COLUMNS = {
    "tool_readiness": ["snapshot_date", "site_code", "tool_id", "tool_code",
                       "tool_type", "process_area", "criticality",
                       "parts_risk_score", "training_risk_score",
                       "readiness_score", "risk_band", "top_reason"],
    "inventory_health": ["snapshot_date", "site_code", "part_id",
                         "part_number", "description", "lead_time_days",
                         "is_consumable", "on_hand_qty", "reorder_point",
                         "safety_stock", "unit_cost", "avg_daily_usage_90d",
                         "days_of_supply", "last_consumption_date",
                         "consumption_qty_180d", "last_counted_date",
                         "is_stockout", "is_below_reorder", "is_excess",
                         "excess_value_usd"],
    "parts_consumption": ["snapshot_date", "site_code", "part_id",
                          "part_number", "description", "is_consumable",
                          "month", "consumption_qty", "avg_daily_usage_90d",
                          "mom_change_qty", "mom_change_pct"],
    "training_coverage": ["snapshot_date", "site_code", "tool_type", "shift",
                          "certified_techs", "target_techs", "coverage_ratio",
                          "certs_expiring_90d", "is_single_point_of_failure"],
    "training_throughput": ["snapshot_date", "site_code", "month",
                            "enrolled_count", "completed_count",
                            "completion_rate", "median_days_to_certify",
                            "stale_in_progress_count"],
}

KEY_COLUMNS = {
    "tool_readiness": ["tool_id", "tool_code", "site_code"],
    "inventory_health": ["site_code", "part_number"],
    "parts_consumption": ["site_code", "part_number", "month"],
    "training_coverage": ["site_code", "tool_type", "shift"],
    "training_throughput": ["site_code", "month"],
}


@pytest.fixture()
def export_dir(engine, tmp_path_factory, monkeypatch):
    monkeypatch.setenv("FABTOOL_DB_SCHEMA", SCHEMA)
    out = tmp_path_factory.mktemp("exports")
    export_reports.main(["--site", "ALL", "--as-of", AS_OF,
                         "--out", str(out)])
    return out


def test_all_csvs_written_with_snapshot(export_dir):
    tableau = export_dir / "tableau"
    for name, expected_cols in CSV_COLUMNS.items():
        path = tableau / f"{name}.csv"
        assert path.exists(), f"missing {path}"
        df = pd.read_csv(path)
        assert list(df.columns) == expected_cols
        assert len(df) > 0
        assert (df["snapshot_date"] == AS_OF).all()


def test_csvs_have_no_null_keys(export_dir):
    tableau = export_dir / "tableau"
    for name, keys in KEY_COLUMNS.items():
        df = pd.read_csv(tableau / f"{name}.csv")
        for col in keys:
            assert df[col].notna().all(), f"null {col} in {name}.csv"


def test_fixture_values_flow_through_consumption_csv(export_dir):
    df = pd.read_csv(export_dir / "tableau" / "parts_consumption.csv")
    # (site 1, PART-001, 2026-01-01): 9 units over the trailing 3 months
    row = df[(df["site_code"] == "TST") & (df["part_number"] == "PART-001")
             & (df["month"] == "2026-01-01")].iloc[0]
    assert row["consumption_qty"] == 3
    assert row["avg_daily_usage_90d"] == pytest.approx(0.1, abs=1e-9)


def test_workbook_sheets_and_reorder_rows(export_dir):
    wb = openpyxl.load_workbook(export_dir / "FabToolReadiness_Tracker.xlsx")
    assert wb.sheetnames == ["Summary", "Reorder Tracker", "Training Tracker",
                             "At-Risk Tools", "Data Dictionary"]
    # fixture: 2 stockouts + 1 below-reorder row
    rt = wb["Reorder Tracker"]
    assert rt.max_row == 4  # header + 3 data rows
    statuses = {rt.cell(r, 11).value for r in range(2, 5)}
    assert statuses == {"STOCKOUT", "BELOW REORDER"}
    # charts present on the Summary sheet
    assert len(wb["Summary"]._charts) == 2
    # number formats applied, not strings: money + percent cells
    summary = wb["Summary"]
    money = None
    pct = None
    for row in summary.iter_rows():
        for cell in row:
            if cell.value == "Excess inventory value":
                money = summary.cell(cell.row, cell.column + 1)
            if cell.value == "Overall cert coverage (vs 2/shift)":
                pct = summary.cell(cell.row, cell.column + 1)
    assert money is not None and money.number_format.startswith("$")
    assert pct is not None and pct.number_format.endswith("%")
